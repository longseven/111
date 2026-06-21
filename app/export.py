"""把改编结果导出为 Word（.docx）——支持单题与整卷（多题）。

公式策略：优先用 pandoc 把 $...$ LaTeX 转成 Word 原生公式（OMML）——它在
Word 公式编辑器里可编辑，且 MathType 可一键「Word 公式 → MathType」整篇转换。
未装 pandoc 时退回 python-docx，公式以可读纯文本呈现，并在文首给出提示。
"""
from __future__ import annotations

import re
import tempfile
from io import BytesIO
from pathlib import Path
from typing import List, Optional, Tuple

from .schemas import AdaptedVariant, ParsedProblem

# 一个"题组"= (原题解析或None, 该原题的改编变体列表)
Group = Tuple[Optional[ParsedProblem], List[AdaptedVariant]]


def pandoc_available() -> bool:
    try:
        import pypandoc

        pypandoc.get_pandoc_path()
        return True
    except Exception:  # noqa: BLE001
        return False


# ---------- Markdown（OMML 路径） ----------
def _variant_md(v: AdaptedVariant, i: int, level: str = "##") -> List[str]:
    return [
        f"{level} 第 {i} 题（难度：{v.difficulty}）",
        "",
        v.stem,
        "",
        f"**参考答案**：{v.answer}",
        "",
        f"**解析**：{v.solution}",
        "",
        f"**知识点**：{'、'.join(v.knowledge_points)}",
        "",
        f"**改编理由**：{v.adaptation_reason}",
        "",
    ]


def _parsed_md(parsed: Optional[ParsedProblem]) -> List[str]:
    if parsed is None:
        return []
    return [
        f"**原题**：{parsed.problem_text}",
        "",
        f"**知识点**：{'、'.join(parsed.knowledge_points)}　**学段**：{parsed.grade_band}",
        "",
    ]


def build_markdown(
    parsed: Optional[ParsedProblem], variants: List[AdaptedVariant], title: str
) -> str:
    out: List[str] = [f"# {title}", ""] + _parsed_md(parsed)
    for i, v in enumerate(variants, 1):
        out += _variant_md(v, i)
    return "\n".join(out)


def build_markdown_paper(groups: List[Group], title: str) -> str:
    out: List[str] = [f"# {title}", ""]
    for gi, (parsed, variants) in enumerate(groups, 1):
        out += [f"## 原第 {gi} 题改编", ""] + _parsed_md(parsed)
        for vi, v in enumerate(variants, 1):
            out += _variant_md(v, vi, level="###")
    return "\n".join(out)


def _export_with_pandoc(md: str) -> bytes:
    import pypandoc

    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "out.docx"
        pypandoc.convert_text(
            md,
            "docx",
            format="markdown+tex_math_dollars+hard_line_breaks",
            outputfile=str(out),
        )
        return out.read_bytes()


# ---------- 纯文本兜底（无 pandoc） ----------
def strip_latex(text: str) -> str:
    def repl(m: "re.Match[str]") -> str:
        s = m.group(1)
        s = re.sub(r"\\frac\s*\{([^{}]*)\}\s*\{([^{}]*)\}", r"(\1)/(\2)", s)
        s = re.sub(r"\\sqrt\s*\{([^{}]*)\}", r"√(\1)", s)
        for a, b in (
            ("\\times", "×"), ("\\cdot", "·"), ("\\div", "÷"),
            ("\\leq", "≤"), ("\\le", "≤"), ("\\geq", "≥"), ("\\ge", "≥"),
            ("\\neq", "≠"), ("\\pm", "±"), ("\\sim", "~"), ("\\%", "%"),
        ):
            s = s.replace(a, b)
        s = re.sub(r"\\left|\\right", "", s)
        s = s.replace("{", "").replace("}", "")
        s = re.sub(r"\\[a-zA-Z]+", "", s)
        return s.strip()

    return re.sub(r"\$\$?(.+?)\$\$?", repl, text, flags=re.S)


def _docx_add_lines(doc, label: str, text: str) -> None:
    p = doc.add_paragraph()
    if label:
        p.add_run(label).bold = True
    for li, line in enumerate(text.split("\n")):
        if li > 0:
            p.add_run().add_break()
        p.add_run(strip_latex(line))


def _docx_add_variant(doc, v: AdaptedVariant, i: int, level: int) -> None:
    doc.add_heading(f"第 {i} 题（难度：{v.difficulty}）", level=level)
    _docx_add_lines(doc, "", v.stem)
    _docx_add_lines(doc, "参考答案：", v.answer)
    _docx_add_lines(doc, "解析：", v.solution)
    doc.add_paragraph(f"知识点：{'、'.join(v.knowledge_points)}")
    _docx_add_lines(doc, "改编理由：", v.adaptation_reason)


def _docx_add_parsed(doc, parsed: Optional[ParsedProblem]) -> None:
    if parsed is None:
        return
    _docx_add_lines(doc, "原题：", parsed.problem_text)
    doc.add_paragraph(f"知识点：{'、'.join(parsed.knowledge_points)}　学段：{parsed.grade_band}")


def _export_with_docx(groups: List[Group], title: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(
        "提示：在运行环境中安装 pandoc（或 pip 安装 pypandoc-binary）后重新导出，"
        "公式将变为可在 Word/MathType 中编辑的格式。"
    )
    paper = len(groups) > 1
    for gi, (parsed, variants) in enumerate(groups, 1):
        if paper:
            doc.add_heading(f"原第 {gi} 题改编", level=1)
        _docx_add_parsed(doc, parsed)
        for vi, v in enumerate(variants, 1):
            _docx_add_variant(doc, v, vi, level=2 if paper else 1)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- 图片公式模式 ----------
def _normalize_latex(s: str) -> str:
    s = re.sub(r"\\le(?![a-zA-Z])", r"\\leq", s)
    s = re.sub(r"\\ge(?![a-zA-Z])", r"\\geq", s)
    s = re.sub(r"\\ne(?![a-zA-Z])", r"\\neq", s)
    return s


def _render_math_png(latex: str) -> Optional[bytes]:
    try:
        from matplotlib.backends.backend_agg import FigureCanvasAgg
        from matplotlib.figure import Figure

        fig = Figure()
        FigureCanvasAgg(fig)
        fig.text(0.0, 0.0, f"${_normalize_latex(latex)}$", fontsize=12)
        buf = BytesIO()
        fig.savefig(buf, format="png", dpi=150, bbox_inches="tight", pad_inches=0.03)
        return buf.getvalue()
    except Exception:  # noqa: BLE001
        return None


def _split_segments(text: str):
    parts = re.split(r"(\$\$.+?\$\$|\$.+?\$)", text, flags=re.S)
    segs = []
    for p in parts:
        if not p:
            continue
        if p.startswith("$$") and p.endswith("$$"):
            segs.append(("math", p[2:-2]))
        elif p.startswith("$") and p.endswith("$"):
            segs.append(("math", p[1:-1]))
        else:
            segs.append(("text", p))
    return segs


def _add_rich_paragraph(doc, label: str, text: str) -> None:
    from PIL import Image
    from docx.shared import Inches

    p = doc.add_paragraph()
    if label:
        p.add_run(label).bold = True
    for li, line in enumerate(text.split("\n")):
        if li > 0:
            p.add_run().add_break()
        for kind, content in _split_segments(line):
            if kind == "text":
                p.add_run(content)
                continue
            png = _render_math_png(content)
            if png:
                w, h = Image.open(BytesIO(png)).size
                p.add_run().add_picture(BytesIO(png), height=Inches(h / 200))
            else:
                p.add_run(strip_latex("$" + content + "$"))


def _img_add_variant(doc, v: AdaptedVariant, i: int, level: int) -> None:
    doc.add_heading(f"第 {i} 题（难度：{v.difficulty}）", level=level)
    _add_rich_paragraph(doc, "", v.stem)
    _add_rich_paragraph(doc, "参考答案：", v.answer)
    _add_rich_paragraph(doc, "解析：", v.solution)
    doc.add_paragraph(f"知识点：{'、'.join(v.knowledge_points)}")
    _add_rich_paragraph(doc, "改编理由：", v.adaptation_reason)


def _export_with_images(groups: List[Group], title: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=0)
    paper = len(groups) > 1
    for gi, (parsed, variants) in enumerate(groups, 1):
        if paper:
            doc.add_heading(f"原第 {gi} 题改编", level=1)
        if parsed is not None:
            _add_rich_paragraph(doc, "原题：", parsed.problem_text)
            doc.add_paragraph(
                f"知识点：{'、'.join(parsed.knowledge_points)}　学段：{parsed.grade_band}"
            )
        for vi, v in enumerate(variants, 1):
            _img_add_variant(doc, v, vi, level=2 if paper else 1)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- 统一入口 ----------
def _normalize_groups(groups: List[Group]) -> List[Group]:
    """导出前对题干/选项分行规整，兼容未分行的数据。"""
    from .claude_service import format_choices

    out: List[Group] = []
    for parsed, variants in groups:
        if parsed is not None:
            parsed = parsed.model_copy(
                update={"problem_text": format_choices(parsed.problem_text)}
            )
        variants = [v.model_copy(update={"stem": format_choices(v.stem)}) for v in variants]
        out.append((parsed, variants))
    return out


def _render(groups: List[Group], title: str, formula_mode: str, paper: bool) -> bytes:
    groups = _normalize_groups(groups)
    if formula_mode == "image":
        return _export_with_images(groups, title)
    if pandoc_available():
        try:
            md = build_markdown_paper(groups, title) if paper else build_markdown(
                groups[0][0], groups[0][1], title
            )
            return _export_with_pandoc(md)
        except Exception:  # noqa: BLE001
            pass
    return _export_with_docx(groups, title)


def export_docx(
    parsed: Optional[ParsedProblem],
    variants: List[AdaptedVariant],
    title: str = "改编题目",
    formula_mode: str = "omml",
) -> bytes:
    """单题导出。"""
    return _render([(parsed, variants)], title, formula_mode, paper=False)


def export_paper_docx(
    groups: List[Group], title: str = "改编试卷", formula_mode: str = "omml"
) -> bytes:
    """整卷导出：groups 为多个 (原题, 变体列表) 题组。"""
    return _render(groups, title, formula_mode, paper=True)
